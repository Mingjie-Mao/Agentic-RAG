"""Small, explicit regression files covering structure, damage and formulas."""

from io import BytesIO
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, PageBreak
from reportlab.pdfgen import canvas

ROOT = Path("fixtures/s2")
ROOT.mkdir(parents=True, exist_ok=True)
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=11, leading=17)
rows = []


def save(name, data, facts=(), error=False):
    path = ROOT / name
    path.write_bytes(data)
    rows.append(
        {
            "path": str(path),
            "format": path.suffix[1:],
            "sha256": hashlib.sha256(data).hexdigest(),
            "facts": list(facts),
            "expect_error": error,
            "provenance": "自建格式回归；不进入问答质量基准",
        }
    )


for name, text, facts, error in [
    ("headings.md", "# 采购\n\n## 审批\n\n采购满 800 澳元需经理审批。\n", ["800"], False),
    ("table.md", "# 配额\n\n| 项目 | 上限 |\n| --- | --- |\n| 批处理 | 40 条 |\n", ["40"], False),
    ("code.md", "# 示例\n\n```python\nretry_limit = 7\nprint('a|b')\n```\n", ["retry_limit = 7"], False),
    ("long.md", "# 长段落\n\n" + "长文定位样本，不得丢失中间的文字。" * 130, ["不得丢失"], False),
    ("empty.md", "   \n", [], True),
]:
    save(name, text.encode(), facts, error)
save("binary.md", b"\xff\x00", error=True)


def pdf_bytes(story):
    output = BytesIO()
    SimpleDocTemplate(output).build(story)
    return output.getvalue()


plain = pdf_bytes([Paragraph("备份说明：恢复目标为 45 分钟。", style)])
save("text.pdf", plain, ["45"])
save(
    "pages.pdf",
    pdf_bytes(
        [
            Paragraph("第一页：批次号 ALPHA-19。", style),
            PageBreak(),
            Paragraph("第二页：归档周期为 73 天。", style),
        ]
    ),
    ["ALPHA-19", "73"],
)
table = Table(
    [["区域", "配额（次/日）"]] + [[f"区域 {i}", str(50 + i)] for i in range(1, 55)], repeatRows=1
)
table.setStyle(
    [
        ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, (0, 0, 0)),
    ]
)
save("multipage-table.pdf", pdf_bytes([Paragraph("地区配额表", style), table]), ["51", "104"])
blank = BytesIO()
c = canvas.Canvas(blank)
c.rect(20, 20, 100, 100)
c.showPage()
c.save()
save("no-text.pdf", blank.getvalue(), error=True)
save("broken.pdf", b"%PDF-1.7\nbroken", error=True)
writer = PdfWriter()
writer.append(PdfReader(BytesIO(plain)))
writer.encrypt("sample-password")
encrypted = BytesIO()
writer.write(encrypted)
save("encrypted.pdf", encrypted.getvalue(), error=True)

for name, variant in [
    ("headings.docx", "headings"),
    ("table.docx", "table"),
    ("long.docx", "long"),
    ("empty.docx", "empty"),
]:
    doc = Document()
    if variant == "headings":
        doc.add_heading("设备申请", 0)
        doc.add_heading("配件", 1)
        doc.add_paragraph("显示器需在 4 个工作日内登记。")
    elif variant == "table":
        doc.add_heading("值班表", 1)
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "班次"
        table.rows[0].cells[1].text = "响应分钟"
        row = table.add_row()
        row.cells[0].text = "夜班"
        row.cells[1].text = "35"
    elif variant == "long":
        doc.add_heading("运行手册", 1)
        doc.add_paragraph("不要把文档页码伪造为稳定页码。" * 100)
    output = BytesIO()
    doc.save(output)
    save(
        name,
        output.getvalue(),
        ["4"] if variant == "headings" else ["35"] if variant == "table" else [],
        variant == "empty",
    )
save("broken.docx", b"not-office", error=True)
doc = Document()
doc.add_heading("Mixed 中文 API", 1)
doc.add_paragraph("HTTP E419 means session rotated；等待 8 秒。")
out = BytesIO()
doc.save(out)
save("mixed.docx", out.getvalue(), ["E419", "8"])

for name, variant in [
    ("units.xlsx", "units"),
    ("merged.xlsx", "merged"),
    ("formula-missing.xlsx", "missing"),
    ("formula-cached.xlsx", "cached"),
    ("empty.xlsx", "empty"),
]:
    book = Workbook()
    sheet = book.active
    sheet.title = "费用规则"
    if variant != "empty":
        sheet.append(["项目", "金额（AUD）", "上限"])
        sheet.append(["交通", 27, 130])
        sheet["B2"].number_format = '0.00 "AUD"'
    if variant == "merged":
        sheet.merge_cells("A4:C4")
        sheet["A4"] = "所有金额以澳元计价"
    if variant in {"missing", "cached"}:
        sheet["D1"] = "合计"
        sheet["D2"] = "=B2+C2"
    out = BytesIO()
    book.save(out)
    data = out.getvalue()
    if variant == "cached":
        replaced = BytesIO()
        with ZipFile(BytesIO(data)) as source, ZipFile(replaced, "w", ZIP_DEFLATED) as dest:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    content = content.replace(b"<f>B2+C2</f><v></v>", b"<f>B2+C2</f><v>157</v>")
                dest.writestr(info, content)
        data = replaced.getvalue()
    save(name, data, ["27", "130"] if variant != "empty" else [], variant == "empty")
save("broken.xlsx", b"broken", error=True)
(ROOT / "manifest.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n")
print(f"Created {len(rows)} format regression fixtures")
