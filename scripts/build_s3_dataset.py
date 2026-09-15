"""Author a versioned 120-document / 200-question corpus from explicit business rules."""

import hashlib
import json
from pathlib import Path

from s3_content import GENRES, document_sections

from docx import Document
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph

ROOT = Path("fixtures/s3")
ROOT.mkdir(parents=True, exist_ok=True)
if (ROOT / "freeze.json").exists():
    raise SystemExit("Dataset is frozen; create an explicitly new version rather than overwriting gold")
(ROOT / "documents").mkdir(exist_ok=True)
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=11, leading=17)
manifest = json.loads(Path("fixtures/manifest.json").read_text())
questions = json.loads(Path("fixtures/questions.json").read_text())["questions"]
for q in questions:
    q.update(
        split="development",
        family=next(
            (d["family"] for d in manifest if d["id"] in ["seed-" + key for key in q["sources"]]),
            "s0-negative",
        ),
        kind="conflict"
        if q["expected"] == "conflict"
        else "unanswerable"
        if not q["sources"]
        else "lookup",
        source_ids=["seed-" + key for key in q["sources"]],
        review={"status": "requires_human_review", "author": "Codex", "source_check": "S0 source audit"},
    )
for q in questions:
    if q["id"] in {"Q18", "Q20"}:
        q.update(
            kind="permission",
            hidden_source_ids=["seed-release-code" if q["id"] == "Q18" else "seed-hc-gateway"],
            hidden_facts=q["forbidden"],
        )
topics = json.loads(Path("fixtures/s3_topics.json").read_text())["topics"]
editorial = json.loads(Path("fixtures/s3_editorial.json").read_text())
query_variants = json.loads(Path("fixtures/s3_query_variants.json").read_text())
new_questions = []
for topic_index, topic in enumerate(topics):
    family_docs = []
    for item_index, item in enumerate(topic["items"]):
        name, field_a, value_a, field_b, value_b = item
        key = f"eval-{topic['id']}-{item_index + 1}"
        title = f"星桥{name}{GENRES[item_index]}"
        version = "2026-09"
        sections = document_sections(item, editorial[topic["id"]][item_index], item_index)
        paragraphs = [f"{title}（{version}）"]
        for heading, lines in sections:
            paragraphs.extend([heading, *lines])
        extension = ["md", "pdf", "docx", "xlsx"][(topic_index + item_index) % 4]
        path = ROOT / "documents" / (key + "." + extension)
        if extension == "md":
            path.write_text(
                "# "
                + paragraphs[0]
                + "\n\n"
                + "\n\n".join(
                    "## " + heading + "\n\n" + "\n\n".join(lines) for heading, lines in sections
                )
                + "\n"
            )
        elif extension == "pdf":
            SimpleDocTemplate(str(path)).build([Paragraph(p, style) for p in paragraphs])
        elif extension == "docx":
            doc = Document()
            doc.add_heading(paragraphs[0], 0)
            for heading, lines in sections:
                doc.add_heading(heading, 1)
                for text in lines:
                    doc.add_paragraph(text)
            doc.save(path)
        else:
            book = Workbook()
            sheet = book.active
            sheet.title = "执行规则"
            sheet.append(["事项", "约束", "规定"])
            sheet.append([name, field_a, value_a])
            sheet.append([name, field_b, value_b])
            sheet.append(["适用范围", "组织与生效日", "星桥软件，2026-09-01"])
            sheet.append(["例外", name + "例外申请", "当周负责人确认"])
            sheet.append(["来源", "虚构资料", "本项目自建"])
            notes = book.create_sheet("处理记录")
            notes.append(["环节", "说明"])
            for heading, lines in sections:
                for text in lines:
                    notes.append([heading, text])
            book.save(path)
        private = item_index == 2
        row = {
            "id": key,
            "source_key": key,
            "tenant_id": "xingqiao",
            "owner_id": "xq-admin",
            "path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "title": title,
            "family": topic["id"],
            "split": topic["split"],
            "tenant_public": not private,
            "groups": [topic["group"]] if private else [],
            "effective_from": "2026-09-01",
            "effective_to": None,
            "version": "s3-authored-v2",
            "provenance": "Synthetic: s3_topics.json + s3_editorial.json; score-independent authoring",
            "genre": GENRES[item_index],
        }
        manifest.append(row)
        family_docs.append((row, item))
        user = "xq-engineer" if topic["group"] == "engineering" else "xq-support"
        for field_index, (field, value) in enumerate([(field_a, value_a), (field_b, value_b)]):
            n = topic_index * 10 + item_index * 2 + field_index + 21
            q = {
                "id": f"Q{n:03}",
                "user": user,
                "question": ("星桥：" + query_variants[topic["id"]][item_index])
                if field_index == 0
                else f"星桥{name}的{field}是多少？",
                "expected": "answered",
                "facts": [value],
                "source_ids": [key],
                "split": topic["split"],
                "family": topic["id"],
                "kind": "lookup",
                "review": {
                    "status": "requires_human_review",
                    "author": "Codex",
                    "source_check": "rule-table + parsed-source validation pending",
                },
            }
            if item_index == 2 and field_index == 1:
                q.update(
                    user="xq-support" if user == "xq-engineer" else "xq-engineer",
                    expected="insufficient_evidence",
                    facts=[],
                    source_ids=[],
                    hidden_source_ids=[key],
                    hidden_facts=[value],
                    kind="permission",
                )
            if item_index == 3 and field_index == 1:
                q.update(
                    question=f"星桥{name}在南极分部的专用补贴金额是多少？",
                    expected="insufficient_evidence",
                    facts=[],
                    source_ids=[],
                    kind="unanswerable",
                )
            if item_index == 4 and field_index == 1:
                other, other_item = family_docs[3]
                q.update(
                    question=f"请对照星桥的{other_item[0]}和{name}：前者的{other_item[1]}、后者的{field_b}分别是多少？",
                    facts=[other_item[2], value_b],
                    source_ids=[other["id"], key],
                    kind="multi_document",
                )
            new_questions.append(q)
questions.extend(new_questions)
assert len(manifest) == 120 and len(questions) == 200
assert sum(q["split"] == "development" for q in questions) == 120
assert sum(q["split"] == "holdout" for q in questions) == 80
train = {q["family"] for q in questions if q["split"] == "development"}
held = {q["family"] for q in questions if q["split"] == "holdout"}
assert not train & held
# S0 source families, including the two travel versions, are always development;
# the unused invoice document remains within the held-out billing family.
for row in manifest:
    if row["id"].startswith("seed-"):
        row["owner_id"] = "hc-admin" if row["tenant_id"] == "haichuan" else "xq-admin"
(ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
(ROOT / "questions.json").write_text(json.dumps(questions, ensure_ascii=False, indent=2) + "\n")
(ROOT / "split.json").write_text(
    json.dumps(
        {
            "version": "s3-authored-v2",
            "seed": 42,
            "documents": 120,
            "questions": 200,
            "development": 120,
            "holdout": 80,
            "method": "author-assigned topic groups before query evaluation; all versions stay together",
            "development_topics": sorted(train),
            "holdout_topics": sorted(held),
            "review_state": "source-auditable draft; independent human review not yet performed",
            "limitations": "synthetic business rules and authored incident notes; five genres with shared procedural language; not independently collected enterprise data",
        },
        ensure_ascii=False,
        indent=2,
    )
    + "\n"
)
print(
    "Created 120 documents / 200 questions; development 120, holdout 80; human review remains explicit"
)
