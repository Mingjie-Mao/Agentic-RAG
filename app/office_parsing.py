"""Office XML readers with original paragraph/cell coordinates; never evaluate formulas."""

from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.parsing import ParseError, Passage


def validate_archive(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                raise ParseError("Office 文件解压后超过 50 MB 或条目过多")
            if any(
                PurePosixPath(e.filename).is_absolute() or ".." in PurePosixPath(e.filename).parts
                for e in entries
            ):
                raise ParseError("Office 文件包含不合法的内部路径")
            if len({e.filename for e in entries}) != len(entries):
                raise ParseError("Office 文件包含重复条目")
    except BadZipFile as exc:
        raise ParseError("Office 文件不是有效的 DOCX/XLSX 压缩包") from exc


def parse_docx(data):
    validate_archive(data)
    try:
        document = WordDocument(BytesIO(data))
        result, headings, paragraph_number, table_number = [], [], 0, 0
        for element in document.element.body.iterchildren():
            if element.tag.endswith("}p"):
                paragraph_number += 1
                paragraph = Paragraph(element, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                style = paragraph.style.name if paragraph.style else ""
                level = (
                    int(style.split()[-1])
                    if style.startswith("Heading ") and style.split()[-1].isdigit()
                    else 0
                )
                if level:
                    headings = headings[: level - 1] + [text]
                result.append(
                    Passage(
                        text,
                        {
                            "kind": "docx",
                            "block_type": "heading" if level else "paragraph",
                            "heading_path": list(headings),
                            "paragraph": paragraph_number,
                            "label": f"段落 {paragraph_number}",
                        },
                    )
                )
            elif element.tag.endswith("}tbl"):
                table_number += 1
                table = Table(element, document)
                headers = [cell.text.strip() for cell in table.rows[0].cells] if table.rows else []
                for row_number, row in enumerate(table.rows, 1):
                    values = [cell.text.strip() for cell in row.cells]
                    if not any(values):
                        continue
                    text = " | ".join(values)
                    if row_number > 1:
                        text = "表头：" + " | ".join(headers) + "\n" + text
                    result.append(
                        Passage(
                            text,
                            {
                                "kind": "docx",
                                "block_type": "table_row",
                                "heading_path": list(headings),
                                "table": table_number,
                                "row": row_number,
                                "label": f"表 {table_number} · 第 {row_number} 行",
                                "headers": headers,
                            },
                        )
                    )
        if not result:
            raise ParseError("Word 文件没有可检索文字")
        return result
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("Word 文件损坏或结构无法读取") from exc


def parse_xlsx(data):
    validate_archive(data)
    formulas = values = None
    try:
        formulas = load_workbook(BytesIO(data), data_only=False, read_only=False, keep_links=False)
        values = load_workbook(BytesIO(data), data_only=True, read_only=False, keep_links=False)
        if len(formulas.worksheets) > 30:
            raise ParseError("工作簿最多支持 30 个工作表")
        result = []
        for sheet in formulas.worksheets:
            if sheet.sheet_state != "visible":
                continue
            if sheet.max_row * sheet.max_column > 200000:
                raise ParseError("工作表范围过大，最多支持 200000 个单元格")
            cached = values[sheet.title]
            merged = list(sheet.merged_cells.ranges)
            headers = None
            for row in sheet.iter_rows():
                cells, displays = [], []
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    formula = value if cell.data_type == "f" else None
                    cached_value = cached[cell.coordinate].value if formula else None
                    display = str(value)
                    if formula:
                        display = f"公式 {formula}；" + (
                            f"缓存结果 {cached_value}"
                            if cached_value is not None
                            else "结果未缓存，不能作为已知数值"
                        )
                    cells.append(
                        {
                            "coordinate": cell.coordinate,
                            "value": None if formula else str(value),
                            "formula": formula,
                            "cached_value": str(cached_value) if cached_value is not None else None,
                            "cache_status": "missing"
                            if formula and cached_value is None
                            else "available"
                            if formula
                            else "not_formula",
                            "number_format": cell.number_format,
                            "merged_range": next((str(m) for m in merged if cell.coordinate in m), None),
                        }
                    )
                    displays.append(f"{cell.coordinate}: {display}")
                if not cells:
                    continue
                is_header = headers is None
                if is_header:
                    headers = [c["value"] or c["formula"] for c in cells]
                start, end = cells[0]["coordinate"], cells[-1]["coordinate"]
                text = "\n".join(displays)
                if not is_header:
                    text = "表头：" + " | ".join(headers) + "\n" + text
                row_number = row[0].row
                result.append(
                    Passage(
                        text,
                        {
                            "kind": "xlsx",
                            "block_type": "table_row",
                            "sheet": sheet.title,
                            "range": f"{start}:{end}",
                            "row": row_number,
                            "heading_path": [sheet.title],
                            "headers": headers,
                            "cells": cells,
                            "label": f"{sheet.title} · {start}:{end}",
                        },
                    )
                )
        if not result:
            raise ParseError("Excel 可见工作表没有可检索内容")
        return result
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("Excel 文件损坏或结构无法读取") from exc
    finally:
        if formulas:
            formulas.close()
        if values:
            values.close()
